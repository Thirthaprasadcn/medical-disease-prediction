"""Generate academic PPT for Medical Imaging Report / Disease Prediction project."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "Medical_Disease_Prediction_Presentation.pptx"

# Clinical light theme
INK = RGBColor(0x0F, 0x1A, 0x22)
MUTE = RGBColor(0x3A, 0x4B, 0x56)
ACCENT = RGBColor(0x0C, 0x7C, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF3, 0xF5, 0xF7)
LINE = RGBColor(0xD0, 0xD8, 0xDF)
SOFT = RGBColor(0xE8, 0xF5, 0xF1)


def _set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_bg(slide, prs, color=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        prs.slide_width,
        prs.slide_height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _accent_bar(slide, prs):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def _title_block(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, size=28, bold=True, color=INK)
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        _set_run(r2, size=14, color=MUTE)


def _bullets(slide, lines: list[str], top=1.15, left=0.55, width=12.2, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "•  " + line
        _set_run(run, size=size, color=INK)


def _footer(slide, prs, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(6.95), Inches(12.2), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Medical Disease Prediction  |  Research prototype — not a medical device  |  {page}/{total}"
    _set_run(run, size=10, color=MUTE)


def new_content_slide(prs, title, lines, page, total, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _add_bg(slide, prs)
    _accent_bar(slide, prs)
    _title_block(slide, title, subtitle)
    # thin line under title
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    _bullets(slide, lines)
    _footer(slide, prs, page, total)
    return slide


def title_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, prs, WHITE)
    # top accent band
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    band.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Medical Imaging Report"
    _set_run(r, size=36, bold=True, color=INK)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "AI-Based Multi-Modality Disease Prediction"
    _set_run(r2, size=24, color=ACCENT)
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = "Brain MRI  ·  Chest X-ray  ·  Chest CT"
    _set_run(r3, size=16, color=MUTE)

    meta = slide.shapes.add_textbox(Inches(0.8), Inches(4.6), Inches(11.5), Inches(1.4))
    tfm = meta.text_frame
    for i, line in enumerate(
        [
            "Presented by: Thirthaprasad C N",
            "Email: thirthaprasadcn@gmail.com",
            "Project: medical-disease-prediction",
        ]
    ):
        p = tfm.paragraphs[0] if i == 0 else tfm.add_paragraph()
        run = p.add_run()
        run.text = line
        _set_run(run, size=16, color=INK)
    _footer(slide, prs, 1, total)


def table_slide(prs, title, headers, rows, page, total, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, prs)
    _accent_bar(slide, prs)
    _title_block(slide, title)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()

    cols = len(headers)
    table_shape = slide.shapes.add_table(len(rows) + 1, cols, Inches(0.55), Inches(1.25), Inches(12.2), Inches(0.45 * (len(rows) + 1)))
    table = table_shape.table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                _set_run(run, size=13, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _set_run(run, size=12, color=INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else SOFT

    if note:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(5.8), Inches(12.2), Inches(0.8))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = note
        _set_run(run, size=13, color=MUTE)
    _footer(slide, prs, page, total)


def flow_slide(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, prs)
    _accent_bar(slide, prs)
    _title_block(slide, "Flow Diagram", "End-to-end imaging pipeline")
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()

    steps = [
        (0.55, "1. Upload", "PNG / JPG / DICOM"),
        (3.0, "2. Preprocess", "Resize · Denoise · CLAHE"),
        (5.45, "3. Route", "Brain / XR / CT"),
        (7.9, "4. Infer", "Disease models"),
        (10.35, "5. Report", "Review · FHIR"),
    ]
    for x, title, sub in steps:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(2.2), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = ACCENT
        tf = card.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = title
        _set_run(r, size=14, bold=True, color=ACCENT)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        _set_run(r2, size=11, color=MUTE)

    # branch boxes
    branches = [
        (1.5, "Brain MRI\nTumor CNN + Seg"),
        (5.3, "Chest X-ray\nPneumonia CNN"),
        (9.1, "Chest CT\nNodule CNN"),
    ]
    for x, text in branches:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.5), Inches(2.8), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = SOFT
        card.line.color.rgb = ACCENT
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        _set_run(r, size=13, bold=True, color=INK)

    note = slide.shapes.add_textbox(Inches(0.55), Inches(5.1), Inches(12.2), Inches(1.4))
    tf = note.text_frame
    tf.word_wrap = True
    lines = [
        "Routing rule: Scan type + anatomy detection selects the disease model.",
        "Brain MRI → tumor classification / segmentation · Chest X-ray → pneumonia · Chest CT → nodules.",
        "Output: overlay findings, disease assessment, clinician accept/reject, FHIR export.",
    ]
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "•  " + t
        _set_run(run, size=14, color=INK)
    _footer(slide, prs, page, total)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides_plan = []  # for counting; we build manually with known total
    total = 13

    title_slide(prs, total)

    new_content_slide(
        prs,
        "Contents",
        [
            "Abstract",
            "Introduction",
            "Problem statement",
            "Literature survey",
            "Objective of the proposed project",
            "Proposed methodology",
            "Flow diagram",
            "Outcome of the project",
            "Results",
            "Conclusion",
            "References that support the study",
        ],
        page=2,
        total=total,
    )

    new_content_slide(
        prs,
        "Abstract",
        [
            "This project presents an AI-assisted medical imaging report system for multi-modality disease prediction.",
            "Pipeline: Ingestion → Preprocessing → Inference → Report generation → Clinician review.",
            "Supports Brain MRI (tumor typing), Chest X-ray (pneumonia), and Chest CT (nodule screening).",
            "Implemented with FastAPI, PyTorch CNNs, DICOM ingestion, FHIR export, and a multi-page web UI.",
            "Research prototype only — not a medical device and not clinically validated.",
        ],
        page=3,
        total=total,
    )

    new_content_slide(
        prs,
        "Introduction",
        [
            "Radiology workloads are increasing; timely reporting remains a challenge.",
            "Different imaging modalities require specialized AI models — one model is not enough.",
            "A practical system must upload studies, route correctly, analyze disease, and produce a reviewable report.",
            "This project builds a local industrial-style console: Home → Upload → Output.",
            "It combines classical preprocessing, trained deep models, and clinical workflow features (review, FHIR, PACS scaffold).",
        ],
        page=4,
        total=total,
    )

    new_content_slide(
        prs,
        "Problem Statement",
        [
            "Manual interpretation is time-consuming and varies across Brain MRI, X-ray, and CT.",
            "Using a single chest model (e.g., pneumonia) on brain MRI produces incorrect disease labels.",
            "PNG/JPG uploads often lack DICOM modality tags, increasing misrouting risk.",
            "Need an anatomy-aware router plus separate trained models per modality.",
            "Need structured outputs: overlays, disease assessment, clinician sign-off, and interoperability (FHIR).",
        ],
        page=5,
        total=total,
    )

    table_slide(
        prs,
        "Literature Survey",
        ["Area", "Approach / resource", "Relevance to project"],
        [
            ["Chest X-ray AI", "PneumoniaMNIST / CNN classifiers", "Pneumonia vs normal path"],
            ["Brain MRI AI", "Tumor MRI classification datasets", "4-class tumor CNN"],
            ["CT nodules", "LIDC-IDRI / nodule benchmarks", "Chest CT nodule path"],
            ["Segmentation", "BraTS / U-Net style models", "Brain tumor region findings"],
            ["Standards", "DICOM & FHIR R4", "Ingest + DiagnosticReport export"],
        ],
        page=6,
        total=total,
        note="Selected public datasets and standards guide the industrial pipeline design.",
    )

    new_content_slide(
        prs,
        "Objective of the Proposed Project",
        [
            "Develop a web-based medical imaging report system for study upload and analysis.",
            "Automatically route each scan to Brain MRI, Chest X-ray, or Chest CT pathways.",
            "Train and integrate modality-specific disease models with reported test metrics.",
            "Generate overlays, structured reports, and FHIR DiagnosticReport export.",
            "Enable clinician review actions (accept/reject findings) and report sign-off.",
        ],
        page=7,
        total=total,
    )

    new_content_slide(
        prs,
        "Proposed Methodology",
        [
            "Frontend: industrial multi-page UI — Home, Upload, Output — with 3D imaging stage.",
            "Backend: FastAPI services for upload, inference, registry, review, and PACS scaffolding.",
            "Preprocessing: resize, denoise, CLAHE contrast enhancement.",
            "Modality router: user scan type + image/filename heuristics → disease model.",
            "Models: Brain tumor CNN, Pneumonia CNN, CT nodule CNN, Brain Mini-UNet segmentation.",
            "Post-process: findings overlay, disease assessment, SQLite review persistence, FHIR bundle.",
        ],
        page=8,
        total=total,
    )

    flow_slide(prs, page=9, total=total)

    new_content_slide(
        prs,
        "Outcome of the Project",
        [
            "Working Medical Imaging Report application with Home → Upload → Output flow.",
            "Four ready disease models displayed in the UI with status and metrics.",
            "Safer routing so brain MRI is not incorrectly labeled using pneumonia models.",
            "Report page with modality, anatomy, disease assessment, findings, and exports.",
            "Repository published: github.com/Thirthaprasadcn/medical-disease-prediction",
            "Regulatory documentation track prepared (intended use, ISO/IEC outlines).",
        ],
        page=10,
        total=total,
    )

    table_slide(
        prs,
        "Results",
        ["Model", "Task", "Dataset / proxy", "Metric"],
        [
            ["Brain tumor CNN", "4-class MRI tumor typing", "Brain tumor MRI set", "Accuracy ≈ 83%"],
            ["Pneumonia CNN", "Pneumonia vs normal", "PneumoniaMNIST", "Accuracy ≈ 87%"],
            ["CT nodule CNN", "Nodule vs no nodule", "NoduleMNIST3D slices", "Accuracy ≈ 86%"],
            ["Brain Mini-UNet", "Tumor region mask", "MRI + weak masks", "Dice ≈ 77%"],
        ],
        page=11,
        total=total,
        note="Metrics are research/demo results on public proxies — not clinical validation.",
    )

    new_content_slide(
        prs,
        "Conclusion",
        [
            "Multi-modality disease prediction requires separate models and correct routing.",
            "The system delivers an end-to-end research workflow from upload to signed report.",
            "Limitations: not FDA/CDSCO cleared; CT/BraTS paths use research proxy training.",
            "Future work: full LIDC/BraTS clinical training, larger hospital datasets, tighter PACS integration.",
            "Overall, the project demonstrates a practical industrial-style imaging AI console.",
        ],
        page=12,
        total=total,
    )

    new_content_slide(
        prs,
        "References",
        [
            "MedMNIST / PneumoniaMNIST / NoduleMNIST3D — public medical imaging benchmarks.",
            "Brain Tumor Classification MRI dataset (public class-folder MRI collections).",
            "LIDC-IDRI — lung nodule CT dataset (TCIA) for clinical CT detection targets.",
            "BraTS — multimodal brain tumor segmentation challenge datasets.",
            "HL7 FHIR R4 — DiagnosticReport / Observation resources.",
            "DICOM standard — medical image storage and transfer.",
            "Project repository: https://github.com/Thirthaprasadcn/medical-disease-prediction",
        ],
        page=13,
        total=total,
        subtitle="Sources that support this study",
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
